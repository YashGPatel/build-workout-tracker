#!/usr/bin/env python3
"""Build a Nippard-style workout tracker (.xlsx) from a program YAML.

    python src/build_tracker.py programs/arm-shoulder.yaml out.xlsx

Output imports cleanly into Google Sheets: effort dropdown, colour-coded
R/W/E cells, frozen exercise column, rotated superset labels, numeric
formats on R/W/Rest so mobile opens the numeric keypad.
"""
import argparse
import datetime
import sys

import yaml
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

WEEK_COLS = ['R', 'W', 'E', 'Rest', 'Notes']
WEEK_WIDTHS = [3.88, 4.38, 3.0, 6.5, 18.38]
NUMFMTS = ['0', '0.0', None, '0', None]        # R, W, E, Rest, Notes

YELLOW = PatternFill('solid', fgColor='FFF2CC')
WHITE = PatternFill('solid', fgColor='FFFFFF')
GREY = PatternFill('solid', fgColor='EFEFEF')
BLUE_DARK = PatternFill('solid', fgColor='406EDF')
BLUE = PatternFill('solid', fgColor='4285F4')
TAPER = PatternFill('solid', fgColor='D9D9D9')

ROBOTO = dict(name='Roboto', sz=10)
thin_grey = Side(style='thin', color='FFEFEFEF')
GREY_BORDER = Border(left=thin_grey, right=thin_grey, top=thin_grey, bottom=thin_grey)
white_side = Side(style='thin', color='FFFFFFFF')
grid_side = Side(style='thin', color='FFD9D9D9')
WHITE_GRID = Border(left=white_side, right=white_side, top=white_side, bottom=white_side)
GREY_GRID = Border(left=grid_side, right=grid_side, top=grid_side, bottom=grid_side)
thick_black = Side(style='thick', color='FF000000')
medium_black = Side(style='medium', color='FF000000')

EFFORT_STYLE = [('Failed', '990000', False),
                ('10 (Max Effort)', 'FF9900', True),
                ('9 (Hard)', 'FFCC00', False),
                ('8 (Medium)', 'F8F848', False),
                ('7 (Easy)', '8FCB20', False)]
CF_COLORS = {'Failed': 'FF990000', '10 (Max Effort)': 'FFFF5500',
             '9 (Hard)': 'FFFF9900', '8 (Medium)': 'FFFFFF00',
             '7 (Easy)': 'FF8FCB20'}


def set_texts(n, reps):
    return [f'Set {i}: {reps}' for i in range(1, n + 1)]


def build(program, out_path, taper_weeks=0, date_weeks=8):
    n_weeks = program.get('weeks', 50)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for sess in program['sessions']:
        ws = wb.create_sheet(sess['name'])
        first_date = sess['first_date']
        if isinstance(first_date, str):
            first_date = datetime.date.fromisoformat(first_date)

        ws.freeze_panes = 'D3'
        for col, width in zip('ABC', (1.75, 3.38, 12.25)):
            ws.column_dimensions[col].width = width
        for w in range(n_weeks):
            for k, width in enumerate(WEEK_WIDTHS):
                ws.column_dimensions[get_column_letter(4 + w * 5 + k)].width = width

        # ---- header rows -------------------------------------------------
        ws.row_dimensions[1].height = 18.0
        ws.row_dimensions[2].height = 17.25
        ws.merge_cells('A1:C1')
        c = ws['A1']
        c.value, c.font = 'Date', Font(**ROBOTO)
        c.alignment = Alignment(horizontal='right', vertical='center', wrapText=True)
        ws.merge_cells('A2:C2')
        c = ws['A2']
        c.value = 'WORKOUTS'
        c.font = Font(b=True, color='FFFFFFFF', **ROBOTO)
        c.fill = BLUE_DARK
        c.alignment = Alignment(horizontal='center', vertical='center', wrapText=True)

        for w in range(n_weeks):
            c0 = 4 + w * 5
            yellow = (w % 2 == 0)
            fill = YELLOW if yellow else WHITE
            grid = WHITE_GRID if yellow else GREY_GRID
            L0, L4 = get_column_letter(c0), get_column_letter(c0 + 4)
            ws.merge_cells(f'{L0}1:{L4}1')
            dc = ws.cell(row=1, column=c0)
            dc.number_format = 'DD.MM.YYYY'
            if w < date_weeks:
                dc.value = first_date + datetime.timedelta(weeks=w)
            dc.alignment = Alignment(horizontal='center', vertical='center')
            dc.font = Font(**ROBOTO)
            for k in range(5):
                cell = ws.cell(row=1, column=c0 + k)
                cell.fill, cell.border = fill, grid
            for k, h in enumerate(WEEK_COLS):
                hc = ws.cell(row=2, column=c0 + k, value=h)
                if yellow:
                    hc.font, hc.fill, hc.border = Font(name='Lato', sz=10, b=True), WHITE, GREY_GRID
                else:
                    hc.font = Font(name='Lato', sz=10, b=True, color='FFFFFFFF')
                    hc.fill, hc.border = BLUE_DARK, Border()
                hc.alignment = Alignment(horizontal='center', vertical='center')

        # ---- exercise blocks ---------------------------------------------
        row, set_rows, prefill = 3, [], []

        def style_name_row(r, cols):
            ws.row_dimensions[r].height = 49.5
            for col in cols:
                cell = ws.cell(row=r, column=col)
                cell.fill, cell.border = GREY, GREY_BORDER
            for w in range(n_weeks):
                c0 = 4 + w * 5
                ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 4)
                for k in range(5):
                    cell = ws.cell(row=r, column=c0 + k)
                    cell.fill, cell.border = GREY, GREY_BORDER

        def style_set_row(r, start=None):
            ws.row_dimensions[r].height = 17.25
            for w in range(n_weeks):
                c0 = 4 + w * 5
                yellow = (w % 2 == 0)
                for k in range(5):
                    cell = ws.cell(row=r, column=c0 + k)
                    cell.fill = YELLOW if yellow else WHITE
                    cell.border = WHITE_GRID if yellow else GREY_GRID
                    if NUMFMTS[k] and not isinstance(start, str):
                        cell.number_format = NUMFMTS[k]
            set_rows.append(r)
            if start is not None:
                prefill.append((r, start))

        for block in sess['blocks']:
            if block['type'] == 'superset':
                start_row = row
                for ex in block['exercises']:
                    for col, val in ((2, ex['code']), (3, ex['name'])):
                        cell = ws.cell(row=row, column=col, value=val)
                        cell.font = Font(b=True, color='FF222222', **ROBOTO)
                        cell.alignment = Alignment(horizontal='center', vertical='center',
                                                   wrapText=True)
                    style_name_row(row, (2, 3))
                    row += 1
                    for txt in set_texts(ex['sets'], ex['reps']):
                        ws.merge_cells(f'B{row}:C{row}')
                        sc = ws.cell(row=row, column=2, value=txt)
                        sc.font = Font(**ROBOTO)
                        sc.fill, sc.border = WHITE, GREY_GRID
                        ws.cell(row=row, column=3).border = GREY_GRID
                        sc.alignment = Alignment(horizontal='center', vertical='center',
                                                 wrapText=True)
                        style_set_row(row, ex.get('start'))
                        row += 1
                end_row = row - 1
                ws.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
                lc = ws.cell(row=start_row, column=1, value=block['label'])
                lc.font = Font(b=True, color='FFFFFFFF', **ROBOTO)
                lc.fill = BLUE
                lc.alignment = Alignment(horizontal='center', vertical='center',
                                         wrapText=True, textRotation=90)
                ws.cell(row=end_row, column=1).border = Border(bottom=thick_black)
            else:
                ws.merge_cells(f'A{row}:C{row}')
                nc = ws.cell(row=row, column=1, value=block['name'])
                nc.font = Font(b=True, color='FF222222', **ROBOTO)
                nc.alignment = Alignment(horizontal='center', vertical='center', wrapText=True)
                style_name_row(row, (1, 2, 3))
                row += 1
                for txt in set_texts(block['sets'], block['reps']):
                    ws.merge_cells(f'A{row}:C{row}')
                    sc = ws.cell(row=row, column=1, value=txt)
                    sc.font = Font(**ROBOTO)
                    for k in (1, 2, 3):
                        ws.cell(row=row, column=k).border = GREY_GRID
                    sc.alignment = Alignment(horizontal='center', vertical='center',
                                             wrapText=True)
                    style_set_row(row, block.get('start'))
                    row += 1
        last_row = row - 1

        for col in range(1, 4 + n_weeks * 5):
            cell = ws.cell(row=last_row, column=col)
            cell.border = Border(left=cell.border.left, right=cell.border.right,
                                 top=cell.border.top, bottom=medium_black)

        # week-1 starting weights
        for r, weight in prefill:
            ws.cell(row=r, column=5, value=weight)

        # optional return taper: grey out set 3 for the first N weeks
        for r in set_rows:
            label = ws.cell(row=r, column=1).value or ws.cell(row=r, column=2).value
            if not (isinstance(label, str) and label.startswith('Set 3')):
                continue
            for w in range(taper_weeks):
                c0 = 4 + w * 5
                for k in range(5):
                    cell = ws.cell(row=r, column=c0 + k)
                    cell.value, cell.fill = None, TAPER
                ws.cell(row=r, column=c0, value='skip').font = Font(
                    name='Roboto', sz=8, italic=True, color='FF808080')

        # alternating superset order banner
        if sess.get('alternate_supersets'):
            flip = sess.get('alternate_start', 'B')
            for w in range(date_weeks):
                c0 = 4 + w * 5
                cell = ws.cell(row=3, column=c0)
                first = flip if w % 2 == 0 else ('A' if flip == 'B' else 'B')
                cell.value = f'{first} FIRST' if first == flip else f'{first} first'
                cell.font = Font(name='Roboto', sz=9, b=(first == flip),
                                 color='FF9C0006' if first == flip else 'FF666666')
                cell.alignment = Alignment(horizontal='center', vertical='center')

        # ---- effort key, dropdown, conditional formatting -----------------
        key_start = 79
        for i, (label, colour, bold) in enumerate(EFFORT_STYLE):
            cell = ws.cell(row=key_start + i, column=3, value=label)
            cell.fill = PatternFill('solid', fgColor=colour)
            cell.font = Font(b=bold, **ROBOTO)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            tb = Side(style='thin', color='FF000000')
            cell.border = Border(left=tb, right=tb, top=tb, bottom=tb)
        legend = ws.cell(row=key_start - 2, column=3,
                         value='R = Reps, W = Weight, E = Effort (pick from dropdown)')
        legend.font = Font(name='Roboto', sz=8)
        legend.alignment = Alignment(horizontal='left', vertical='center')

        dv = DataValidation(type='list', formula1=f'$C${key_start}:$C${key_start + 6}',
                            allowBlank=True, showErrorMessage=False)
        ws.add_data_validation(dv)
        for w in range(n_weeks):
            e_col = get_column_letter(4 + w * 5 + 2)
            for r in set_rows:
                dv.add(f'{e_col}{r}')

        ranges, s, p = [], set_rows[0], set_rows[0]
        for r in set_rows[1:]:
            if r == p + 1:
                p = r
            else:
                ranges.append((s, p))
                s = p = r
        ranges.append((s, p))
        for w in range(n_weeks):
            c0 = 4 + w * 5
            r_col, e_col = get_column_letter(c0), get_column_letter(c0 + 2)
            sqref = ' '.join(f'{r_col}{a}:{e_col}{b}' for a, b in ranges)
            for label, colour in CF_COLORS.items():
                ws.conditional_formatting.add(
                    sqref, FormulaRule(formula=[f'${e_col}{ranges[0][0]}="{label}"'],
                                       fill=PatternFill('solid', bgColor=colour)))

    wb.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('program', help='path to a program YAML')
    ap.add_argument('output', help='output .xlsx path')
    ap.add_argument('--taper-weeks', type=int, default=0,
                    help='grey out the third set for the first N weeks')
    ap.add_argument('--date-weeks', type=int, default=8,
                    help='how many weeks to pre-date')
    args = ap.parse_args()

    with open(args.program) as fh:
        program = yaml.safe_load(fh)
    path = build(program, args.output, args.taper_weeks, args.date_weeks)
    print(f'built {path}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
