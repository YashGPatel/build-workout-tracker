#!/usr/bin/env python3
"""Parse one or more tracker .xlsx files into a tidy CSV.

    python src/extract_history.py history.csv tracker1.xlsx tracker2.xlsx ...
    python src/extract_history.py history.csv old.xlsx --anchor 2023-08-14

Handles both layouts seen in the wild:
  * 5 columns per week  (R, W, E, Rest, Notes)        — current trackers
  * 8 columns per week  (Number of Sets ..., R, W, E) — older Nippard exports

Week headers that read "WEEK 3" instead of a date are placed using --anchor
(the date of week 1). Rows whose R/W cell holds something like "12,8" (a drop
set, or a paired wrist/reverse curl) keep the raw string in `raw_reps`.
"""
import argparse
import csv
import datetime
import re
import sys

import openpyxl

EFFORT = {'Failed': 11, '10 (Max Effort)': 10, '9 (Hard)': 9,
          '8 (Medium)': 8, '7 (Easy)': 7,
          'Max effort - I could not have done any more reps': 10,
          "Failed - I tried to do another rep but couldn't": 11,
          'Hard - I could have done 1 more rep': 9,
          'Medium - I could have done 2 more reps': 8,
          'Easy - I could have done 3 more reps': 7}

SKIP = re.compile(r'^(Set\b|See Tutorial|SUPERSET|DROP SET|https?://|#N/A|'
                  r'WORKOUTS|Date$|R = Reps|[AB] FIRST|[AB] first)', re.I)


def is_exercise_name(v):
    return (isinstance(v, str) and v.strip() and not SKIP.match(v.strip())
            and v.strip() not in EFFORT and len(v.strip()) > 3)


def first_number(v):
    """12 -> 12.0 ; '8+2' -> 10.0 ; '12,8' -> 12.0 ; 'skip' -> None"""
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        nums = [float(n) for n in re.findall(r'\d+\.?\d*', v)]
        if not nums:
            return None
        return sum(nums) if '+' in v else nums[0]
    return None


def parse_week_header(v, anchor):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if isinstance(v, str):
        s = v.strip()
        m = re.match(r'WEEK\s+(\d+)', s, re.I)
        if m and anchor:
            return anchor + datetime.timedelta(weeks=int(m.group(1)) - 1)
        if s.upper().startswith('DELOAD') and anchor:
            return anchor + datetime.timedelta(weeks=10)
        for fmt in ('%d.%m.%Y', '%d.%m.%y', '%d/%m/%Y', '%Y-%m-%d'):
            try:
                return datetime.datetime.strptime(s, fmt).date()
            except ValueError:
                pass
    return None


def extract(path, anchor=None, phase=None):
    rows = []
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        header = [ws.cell(row=2, column=c).value for c in range(1, 45)]
        stride = 8 if 'Number of Sets' in [h for h in header if h] else 5
        r_off, w_off, e_off = (3, 4, 5) if stride == 8 else (0, 1, 2)
        max_row = min(ws.max_row, 60)

        names, current = {}, None
        for r in range(3, max_row + 1):
            for col in (1, 3, 2):
                if is_exercise_name(ws.cell(row=r, column=col).value):
                    current = ws.cell(row=r, column=col).value.strip()
                    break
            names[r] = current

        for c0 in range(4, ws.max_column + 1, stride):
            date = parse_week_header(ws.cell(row=1, column=c0).value, anchor)
            if date is None:
                continue
            for r in range(3, max_row + 1):
                if not names.get(r):
                    continue
                raw_r = ws.cell(row=r, column=c0 + r_off).value
                raw_w = ws.cell(row=r, column=c0 + w_off).value
                eff = ws.cell(row=r, column=c0 + e_off).value
                reps, weight = first_number(raw_r), first_number(raw_w)
                if reps is None and weight is None:
                    continue
                if isinstance(raw_r, str) and raw_r.strip().lower() == 'skip':
                    continue
                rows.append({
                    'date': date.isoformat(),
                    'phase': phase or '',
                    'session': ws.title,
                    'exercise': names[r],
                    'reps': reps,
                    'weight': weight,
                    'effort': EFFORT.get(eff, ''),
                    'raw_reps': raw_r if isinstance(raw_r, str) else '',
                    'raw_weight': raw_w if isinstance(raw_w, str) else '',
                    'source': path.split('/')[-1],
                })
    return rows


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('output', help='output CSV')
    ap.add_argument('trackers', nargs='+', help='tracker .xlsx files')
    ap.add_argument('--anchor', help='date of WEEK 1, for files with no real dates')
    ap.add_argument('--phase', default='', help='label for these files')
    args = ap.parse_args()

    anchor = datetime.date.fromisoformat(args.anchor) if args.anchor else None
    rows = []
    for path in args.trackers:
        rows.extend(extract(path, anchor, args.phase))
    rows.sort(key=lambda r: (r['date'], r['session']))

    fields = ['date', 'phase', 'session', 'exercise', 'reps', 'weight',
              'effort', 'raw_reps', 'raw_weight', 'source']
    with open(args.output, 'w', newline='') as fh:
        writer = csv.DictWriter(fh, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)

    print(f'{len(rows)} records -> {args.output}')
    if rows:
        print(f'range: {rows[0]["date"]} .. {rows[-1]["date"]}')
        print(f'exercises: {len({r["exercise"] for r in rows})}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
